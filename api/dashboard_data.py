"""
LIVE tab data: reads the nightly S3 Excel report the same way the Streamlit
app's LIVE tab does (~/market-dashboard/app.py), rather than querying
DynamoDB per-symbol. See PHASE1_SPEC.md "REVISED ARCHITECTURE" for why.

Ported from app.py: list_dashboard_dates, load_workbook, parse_hud,
parse_compass, parse_grid, _pct, HUD_GROUPS, CLOCK_Q_COLOR. Rewritten to
return plain dicts (no pandas) and extended to read the "Current Date"
column (index 15) that parse_hud never read, for per-ticker stale_days.
"""

from __future__ import annotations

import datetime
import re
from collections import OrderedDict
from io import BytesIO
from typing import Optional

import boto3

REGION = "ap-southeast-1"
BUCKET = "cmon-stage-backend-369568916817-ap-southeast-1-reports"
PREFIX = "Dashboard/"
MODEL_TABLE = "cmon-stage-backend-model-history"

# Ported as-is from app.py:51-120, including CRYPTO (approved deviation #2 --
# the original spec's hud_group_order omitted it, that was an oversight).
HUD_GROUPS: "OrderedDict[str, tuple]" = OrderedDict([
    ("US EQUITIES", (["DJI", "SPX", "IXIC", "RUT", "VIX"], "SPX")),
    ("INDEX ETF", (["DIA", "SPY", "QQQ", "IWM"], "SPX")),
    ("SECTOR ETF", (
        ["XLB", "XLI", "XLY", "XLC", "XLK", "XME", "XLRE", "XLP", "XLU",
         "XLE", "XOP", "XHB", "PBS", "PBJ", "PEJ", "TAN", "ICLN",
         "XLF", "KBE", "KRE", "KIE", "IAI", "XLV", "XHE",
         "IYT", "JETS", "BLOK", "SOCL", "SOXX", "ROBO", "SKYY",
         "FDN", "HACK", "CIBR", "KWEB", "MJ",
         "ARKK", "ARKG", "ARKW", "ARKF", "ARKQ", "IZRL"],
        "SPX",
    )),
    ("US INTEREST RATES", (["US03MY", "US01Y", "US02Y", "US05Y", "US10Y", "US20Y", "US30Y", "MOVE"], None)),
    ("BONDS ETF", (["SHY", "IEF", "TLT", "TMF"], "SPX")),
    ("SPREADS", (["T10Y2Y", "T10Y3M"], None)),
    ("RATES", (["DFEDTARU", "FEDFUNDS", "CPIAUCSL", "GDP", "DRTSCILM"], None)),
    ("COMMODITIES METALS", (
        ["DBC", "USO", "UNG", "GLD", "GDX", "GDXJ", "SLV", "SIL",
         "JJC", "CPER", "JJN", "WOOD", "SLX", "URA"],
        "USCI",
    )),
    ("COMMODITIES CONT.", (
        ["USCI", "USOIL", "NATGAS", "GOLD", "SILVER", "COPPER",
         "NICKEL", "LITHIUM", "SLX", "WOOD", "URANIUM", "COAL"],
        "USCI",
    )),
    ("AGRICULTURAL", (["DBA", "WEAT", "SOYB", "CORN", "RICE", "CANE", "COTTON"], "DBA")),
    ("COUNTRY ETF", (
        ["KWEB", "FXI", "EWJ", "EWZ", "EWT", "EWG", "EWH", "EWI",
         "EWW", "EWU", "PIN", "IDX", "VNM", "EWM", "EIDO", "EPHE",
         "EWY", "EWA", "EWC", "EWS", "EWP", "EWL", "EZA", "INDA"],
        "SPX",
    )),
    ("FOREIGN RATES", (
        ["JP10Y", "CN10Y", "HK10Y", "PH10Y", "EU10Y", "GB10Y",
         "FR10Y", "DE10Y", "IT10Y", "ES10Y", "SG10Y", "KR10Y"],
        None,
    )),
    ("FX", (
        ["USDPHP", "USDJPY", "USDCNY", "USDAUD", "USDEUR", "USDGBP",
         "USDCHF", "USDSGD", "USDKRW", "USDHKD", "USDIDR", "USDINR",
         "USDRUB", "USDTHB", "USDTRY", "DXY", "UUP"],
        "DXY",
    )),
    ("CRYPTO", (["BTC", "ETH", "BITO"], "DXY")),
])

HUD_GROUP_ORDER = list(HUD_GROUPS.keys())

# Ported as-is from app.py:49 / app.py:36-47.
CLOCK_Q_COLOR = {1: "#E87722", 2: "#00C851", 3: "#88DD44", 4: "#FF4444"}
GRID_Q_LABELS = {
    1: "G1 — Goldilocks",
    2: "G2 — Reflation",
    3: "G3 — Inflation",
    4: "G4 — Deflation",
}
COMPASS_Q_LABELS = {
    1: "C1 — Liquidity↑ Credit↓",
    2: "C2 — Liquidity↑ Credit↑",
    3: "C3 — Liquidity↓ Credit↑",
    4: "C4 — Liquidity↓ Credit↓",
}


def _pct(current, ref) -> Optional[float]:
    """Ported as-is from app.py:238."""
    try:
        if ref and ref != 0:
            return (current - ref) / abs(ref) * 100
    except (TypeError, ZeroDivisionError):
        pass
    return None


def list_dashboard_dates() -> list[str]:
    """Ported from app.py:213. Returns ISO date strings, descending."""
    s3 = boto3.client("s3", region_name=REGION)
    paginator = s3.get_paginator("list_objects_v2")
    dates = []
    for page in paginator.paginate(Bucket=BUCKET, Prefix=PREFIX):
        for obj in page.get("Contents", []):
            fname = obj["Key"].rsplit("/", 1)[-1]
            m = re.match(r"dashboard_(\d{4}-\d{2}-\d{2})\.xlsx$", fname)
            if m:
                dates.append(m.group(1))
    return sorted(dates, reverse=True)


def load_workbook(date_str: str):
    """Ported from app.py:230. Returns (openpyxl.Workbook, s3_last_modified: datetime)."""
    import openpyxl

    s3 = boto3.client("s3", region_name=REGION)
    obj = s3.get_object(Bucket=BUCKET, Key=f"{PREFIX}dashboard_{date_str}.xlsx")
    wb = openpyxl.load_workbook(BytesIO(obj["Body"].read()), data_only=True)
    return wb, obj["LastModified"]


def parse_hud(ws) -> list[dict]:
    """
    Adapted from app.py:247 (parse_hud). Returns plain dicts, not a
    DataFrame (approved deviation #4). Extended to read the "Current Date"
    column (row index 15) that the original never read, for stale_days.
    Column layout confirmed empirically against a live workbook (2026-09-04):
      0 Symbol, 2/4/6/8/10/12/14 Reference Value 1D/5D/7D/1M/3M/6M/1Y,
      15 Current Date, 16 Current Value, 18 STD-7D, 19 EMA-7D.
    """
    today = datetime.date.today()
    records = []
    current_sector = "Unknown"
    for row in ws.iter_rows(min_row=2, values_only=True):
        symbol = row[0] if row else None
        if symbol is None:
            continue
        cur_val = row[16] if len(row) > 16 else None
        if cur_val is None:
            current_sector = str(symbol)
            continue
        ref1d = row[2] if len(row) > 2 else None
        ref5d = row[4] if len(row) > 4 else None
        ref7d = row[6] if len(row) > 6 else None
        ref1m = row[8] if len(row) > 8 else None
        ref3m = row[10] if len(row) > 10 else None
        ref6m = row[12] if len(row) > 12 else None
        ref1y = row[14] if len(row) > 14 else None
        cur_date = row[15] if len(row) > 15 else None
        ema7d = row[19] if len(row) > 19 else None

        stale_days = None
        if cur_date:
            try:
                cur_date_parsed = datetime.date.fromisoformat(str(cur_date)[:10])
                stale_days = (today - cur_date_parsed).days
            except ValueError:
                pass

        records.append({
            "symbol": str(symbol),
            "sector": current_sector,
            "current": cur_val,
            "ema_7d": ema7d,
            "sd_7d": row[18] if len(row) > 18 else None,
            "pct_1d": _pct(cur_val, ref1d),
            "pct_5d": _pct(cur_val, ref5d),
            "pct_7d": _pct(cur_val, ref7d),
            "pct_1m": _pct(cur_val, ref1m),
            "pct_3m": _pct(cur_val, ref3m),
            "pct_6m": _pct(cur_val, ref6m),
            "pct_1y": _pct(cur_val, ref1y),
            "as_of": str(cur_date)[:10] if cur_date else None,
            "stale_days": stale_days,
        })

    # de-dupe by symbol, keep first (matches app.py's drop_duplicates(keep="first"))
    seen = set()
    deduped = []
    for r in records:
        if r["symbol"] in seen:
            continue
        seen.add(r["symbol"])
        deduped.append(r)
    return deduped


def _parse_quadrant_block(ws) -> tuple[list[dict], Optional[int], Optional[str]]:
    """
    Shared shape between parse_compass (app.py:288) and parse_grid (app.py:324):
    row[0]=label, quadrant/since embedded in a "Quadrant (since YYYY-MM-DD)"
    row, row[1]=quadrant int. Ref/Cur column positions differ between the two
    sheets (confirmed empirically), so this returns raw rows for the caller
    to interpret, plus the shared quadrant/since parse.
    """
    metrics_rows = []
    quadrant, since = None, None
    for row in ws.iter_rows(values_only=True):
        if row[0] is None:
            continue
        label = str(row[0])
        if label.startswith("Quadrant"):
            try:
                quadrant = int(row[1])
            except (TypeError, ValueError):
                pass
            m = re.search(r"since (\d{4}-\d{2}-\d{2})", label)
            since = m.group(1) if m else None
            break
        if label in ("Symbol", "Metrics"):
            continue
        metrics_rows.append(row)
    return metrics_rows, quadrant, since


def parse_compass(ws) -> tuple[dict, Optional[int], Optional[str]]:
    """
    Adapted from app.py:288. COMPASS sheet columns (confirmed empirically):
    0 Symbol(label), 1 Reference Date, 2 Reference Value, 3 Current Date,
    4 Current Value, 5 % Change, 6 Trend, 7 Rate.
    Returns metrics as a dict keyed by the RAW label from the sheet (e.g.
    "DFEDTARU", "DRTSCILM") -- see build report re: spec's example keys
    (fed_funds/sofr/bank_lending_pct_yoy) not matching what's actually in
    the sheet.
    """
    rows, quadrant, since = _parse_quadrant_block(ws)
    metrics = {}
    for row in rows:
        label = str(row[0])
        ref_val = row[2] if len(row) > 2 else None
        cur_val = row[4] if len(row) > 4 else None
        if cur_val is not None and isinstance(cur_val, (int, float)):
            metrics[label] = {
                "reference_date": str(row[1]) if len(row) > 1 and row[1] else None,
                "reference_value": ref_val,
                "current_date": str(row[3]) if len(row) > 3 and row[3] else None,
                "current_value": cur_val,
                "pct_change": _pct(cur_val, ref_val) if ref_val is not None else None,
            }
    return metrics, quadrant, since


def parse_grid(ws) -> tuple[dict, Optional[int], Optional[str]]:
    """
    Adapted from app.py:324. GRID sheet columns (confirmed empirically):
    0 Metrics(label), 1 Release Date, 2 Previous Value, 3 Current Value,
    4 % Change (sheet computes this itself, no separate current-date column
    on this sheet -- unlike COMPASS).
    """
    rows, quadrant, since = _parse_quadrant_block(ws)
    metrics = {}
    for row in rows:
        label = str(row[0])
        cur = row[3] if len(row) > 3 else None
        if cur is not None and isinstance(cur, (int, float)):
            prev = row[2] if len(row) > 2 else None
            pct_raw = row[4] if len(row) > 4 else None
            metrics[label] = {
                "release_date": str(row[1]) if len(row) > 1 and row[1] else None,
                "previous_value": prev,
                "current_value": cur,
                "pct_change": pct_raw if isinstance(pct_raw, (int, float)) else None,
            }
    return metrics, quadrant, since


def model_history_fallback(model_name: str) -> Optional[int]:
    """
    Approved deviation #3: if a sheet's quadrant parses as None, query
    model-history directly (PK model_name, SK metrics_date) for the latest
    quadrant, rather than the Streamlit-only local-parquet-cache fallback
    (which doesn't apply to a stateless FastAPI request).
    """
    ddb = boto3.client("dynamodb", region_name=REGION)
    resp = ddb.query(
        TableName=MODEL_TABLE,
        KeyConditionExpression="model_name = :m",
        ExpressionAttributeValues={":m": {"S": model_name}},
        ScanIndexForward=False,
        Limit=1,
    )
    items = resp.get("Items", [])
    if not items:
        return None
    q = items[0].get("quadrant", {}).get("N")
    return int(float(q)) if q is not None else None


def build_live_response(date_str: Optional[str] = None) -> dict:
    available = list_dashboard_dates()
    if not available:
        raise RuntimeError("No dashboard reports found in S3")
    resolved_date = date_str if date_str in available else available[0]

    wb, last_modified = load_workbook(resolved_date)
    sheet_names = wb.sheetnames  # [HUD, COMPASS, GRID, CLOCK]

    hud_records = parse_hud(wb[sheet_names[0]])
    compass_metrics, compass_q, compass_since = parse_compass(wb[sheet_names[1]])
    grid_metrics, grid_q, grid_since = parse_grid(wb[sheet_names[2]])

    compass_stale_note = None
    if compass_q is None:
        compass_q = model_history_fallback("compass_US")
        if compass_q is not None:
            compass_stale_note = "Quadrant from live model-history fallback; not present in today's report."

    grid_stale_note = None
    if grid_q is None:
        grid_q = model_history_fallback("grid_US")
        if grid_q is not None:
            grid_stale_note = "Quadrant from live model-history fallback; not present in today's report."

    hud_by_symbol = {r["symbol"]: r for r in hud_records}
    hud_groups = []
    for group_name, (tickers, rs_denom) in HUD_GROUPS.items():
        group_tickers = [hud_by_symbol[t] for t in tickers if t in hud_by_symbol]
        hud_groups.append({
            "name": group_name,
            "default_rs_denom": rs_denom,
            "tickers": group_tickers,
        })

    return {
        "as_of": resolved_date,
        "generated_at": last_modified.astimezone(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "compass": {
            "quadrant": compass_q,
            "label": COMPASS_Q_LABELS.get(compass_q),
            "since": compass_since,
            "stale_note": compass_stale_note,
            "metrics": compass_metrics,
        },
        "grid": {
            "quadrant": grid_q,
            "label": GRID_Q_LABELS.get(grid_q),
            "since": grid_since,
            "stale_note": grid_stale_note,
            "metrics": grid_metrics,
        },
        "hud_groups": hud_groups,
        "hud_group_order": HUD_GROUP_ORDER,
    }
